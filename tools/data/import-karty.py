#!/usr/bin/env python
# Import kart produktu Liveo (xlsx z eksportu Navireo) → tools/data/karty-liveo.json
#
#   python tools/data/import-karty.py "<folder z plikami *.xlsx>"
#
# Źródło: Google Drive, folder „karty produktu dla białegostoku” (udostępniony przez Liveo, 28.10.2026).
# Wymaga: openpyxl. Plik wynikowy czyta tools/data/generate.mjs (bez zależności po stronie Node).
#
# Zasady (uzgodnione z Michałem 28.10.2026):
# - produkt = kolekcja + kod bryły (np. ARIS AS01), wariant = kolor (osobny wiersz/plik w karcie);
# - tylko meble: uchwyty, LED, zasilacze, panele, półki i fronty do szaf, moduły kuchenne → `excluded`;
# - kuchnie tylko jako gotowe zestawy (TIFFANY T-35A/B/C), bez pojedynczych szafek, frontów i blatów;
# - karty nie mają cen detalicznych — nie importujemy żadnych cen (także dropshipping).
import json, re, sys, glob, os, unicodedata
from datetime import date
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '.'
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'karty-liveo.json')

norm = lambda s: ' '.join(str(s).split()) if s is not None else ''

# kolumny karty (nazwy PL po normalizacji białych znaków) → klucz kanoniczny
COLS = {
    'SYMBOL (NAVIREO)': 'symbol', 'NAZWA HANDLOWA (NAVIREO)': 'tradeName', 'RODZAJ ASORTYMENTU': 'kind',
    'KOLORYSTYKA KORPUSU': 'bodyColor', 'KOLORYSTYKA KORPUSU / BALTU / COKOŁU': 'bodyColor',
    'KOLORYSTYKA FRONTU': 'frontColor', 'KOLORYSTYKA FRONTU / UCHWYTU': 'frontColor',
    'TYP (D-drzwi, S-szuflady, P-pojemnik)': 'layout', 'KOD/SYMBOL PRODUCENTA': 'code',
    'WEWNĘTRZNY KOD PRODUCENTA EAN KOMPLETU': 'ean',
    'SZEROKOŚĆ NETTO (cm)': 'w', 'GŁĘBOKOŚĆ NETTO (cm)': 'd', 'WYSOKOŚĆ NETTO (cm)': 'h',
    'DŁUGOŚĆ NETTO - STOŁY, ŁAWY (cm)': 'len', 'WAGA NETTO (kg)': 'netKg', 'WAGA BRUTTO (kg)': 'grossKg',
    'DOSTAWA W 72 GODZINY': 'delivery72',
    'LINK DO INSTRUKCJI': 'instruction', 'LINK DO SZCZEGÓŁOWYCH WYMIARÓW BRYŁY': 'drawing',
    'LINK DO ANIMACJA BRYŁY 180-360*': 'animation', 'ILOŚĆ PACZEK PRODUKTU': 'packages',
    'KOLEKCJA': 'collection', 'STYL': 'style', 'PRZEZNACZENIE': 'purpose', 'MONTAŻ': 'assembly',
    'WYPOSAŻENIE': 'equipment',
    'KORPUS WYKONANIE MATERIAŁ - RODZAJ, GRUBOŚĆ': 'm_body', 'KORPUS WYKOŃCZENIE': 'm_bodyFinish',
    'BLAT WYKONANIE MATERIAŁ - RODZAJ, GRUBOŚĆ': 'm_top', 'BLAT WYKOŃCZENIE': 'm_topFinish', 'BLAT KOLOR/WYBARWIENIE': 'm_topColor',
    'ŚCIANA TYLNA CZĘŚĆ WIDOCZNA WYKOŃCZENIE': 'm_backFinish', 'ŚCIANA TYLNA KOLOR/WYBARWIENIE': 'm_backColor',
    'PÓŁKA PEŁNA WYKONANIE MATERIAŁ RODZAJ I GRUBOŚĆ': 'm_shelf', 'PÓŁKA PEŁNA DOPUSZCZALNE OBCIĄŻENIE (kg)': 'm_shelfLoad',
    'PÓŁKA PEŁNA ILOŚĆ': 'm_shelfQty', 'PÓŁKA SZKLANA RODZAJ I KOLOR SZKŁA': 'm_glassShelf', 'PÓŁKA SZKLANA GRUBOŚĆ (mm)': 'm_glassShelfMm',
    'PÓŁKA SZKLANA DOPUSZCZALNE OBCIĄŻENIE (kg)': 'm_glassShelfLoad', 'PÓŁKA SZKLANA ILOŚĆ': 'm_glassShelfQty',
    'STANDARDOWE WYPOSAŻENIE SZAFY': 'm_wardrobe',
    'FRONT WYKONANIE MATERIAŁ - RODZAJ I GRUBOŚĆ': 'm_front', 'FRONT WYKOŃCZENIE': 'm_frontFinish',
    'FRONT PRZESZKLONY RODZAJ I KOLOR SZKŁA': 'm_glassFront', 'FRONT PRZESZKLONY GRUBOŚĆ SZKŁA (mm)': 'm_glassFrontMm',
    'FRONT Z LUSTREM - WYMIARY LUSTRA (mm)': 'm_mirror', 'UCHWYT MATERIAŁ, ROZMIAR DŁUGOŚĆ/ROZSTAW': 'm_handle',
    'PROWADNICA SZUFLADY RODZAJ': 'm_runners', 'SZUFLADA MATERIAŁ DNA': 'm_drawerBottom', 'SZUFLADA ILOŚĆ': 'm_drawers',
    'STOPKA/NÓŻKA MATERIAŁ': 'm_legs', 'WYSOKOŚĆ NÓŻKI (cm)': 'm_legH', 'STOPKA DODATKOWA': 'm_extraFoot',
    'OŚWIETLENIE STANDARD/OPCJA': 'm_light', 'OŚWIETLENIE RODZAJ / BARWA': 'm_lightType',
    'KONSTRUKCJA RAMY ŁOŻA': 'm_bedFrame', 'POJEMNIK NA POŚCIEL': 'm_bedStorage', 'STELAŻ': 'm_slats',
    'MATERIAŁ WYKONANIA STELAŻA': 'm_slatsMaterial', 'TYP STELAŻA': 'm_slatsType',
    'WYKONANIE BLATU STOŁU/ŁAWY RODZAJ I GRUBOŚĆ': 'm_tableTop', 'ROZKŁADANY': 'm_extendable', 'PROWADNICA STOŁU/ŁAWY': 'm_tableRunner',
    'COKÓŁ DOLNY WYKONANIE MATERIAŁ - RODZAJ I GRUBOŚĆ': 'm_plinth', 'COKÓŁ WYKOŃCZENIE': 'm_plinthFinish',
    'OPIS MARKETINGOWY MEBLA': 'description',
}
for i in range(1, 10):
    COLS[f'LINK DO ZDJĘĆ POJEDYNCZYCH BRYŁ {i}'] = f'photo{i}'
for i in range(1, 6):
    COLS[f'LINK DO ARANŻACJI {i}'] = f'arr{i}'
for i in range(1, 9):
    for pl, k in (('KOD EAN-13 PACZKI', 'ean'), ('SZEROKOŚĆ PACZKI', 'w'), ('WYSOKOŚĆ PACZKI', 'h'), ('GŁĘBOKOŚĆ PACZKI', 'd'), ('WAGA BRUTTO', 'kg')):
        COLS[f'{pl} {i}' + ('' if pl == 'KOD EAN-13 PACZKI' else ' (cm)' if 'PACZKI' in pl else ' (kg)')] = f'p{i}_{k}'

EMPTY = {None, '', '-', '–', '—', 'brak', 'BRAK', 'nie dotyczy', 'NIE DOTYCZY', '0'}

def clean(v):
    if isinstance(v, str):
        v = norm(v)
    return None if v in EMPTY or v == 0 else v

def read_rows(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    header = [norm(x) for x in next(rows)]
    if header and header[1] == 'TRADE NAME (NAVIREO)':
        # VIGO: ta sama kolejność kolumn co szablon PL, nagłówki po angielsku
        header = TEMPLATE[:len(header)] + header[len(TEMPLATE):]
    keys = [COLS.get(h) for h in header]
    for r in rows:
        d = {k: clean(v) for k, v in zip(keys, r) if k}
        sym = d.get('symbol')
        if not isinstance(sym, str) or not d.get('tradeName') or re.fullmatch(r'\d+', sym):
            continue  # pusty wiersz albo wiersz z numerami kolumn
        yield d

# ---------------------------------------------------------------- klasyfikacja
TYPE_RULES = [  # (regex na rodzaj/nazwę, typ strony) — kolejność ma znaczenie
    (r'ZESTAW MEBLI KUCHENNYCH', 'kuchnia'),
    (r'SZAFKA NOCNA', 'szafka_nocna'),
    (r'[ŁL]O[ŻZ]E|[ŁL][ÓO][ŻZ]KO', 'lozko'),
    (r'SZAFKA NA BUTY', 'szafka_na_buty'),
    (r'\bRTV\b', 'rtv'),
    (r'KOMODA', 'komoda'),
    (r'KREDENS', 'kredens'),
    (r'WITRYNA', 'witryna'),
    (r'REGA[ŁL]', 'regal'),
    (r'TOALETKA', 'toaletka'),
    (r'LUSTRO', 'lustro'),
    (r'BIURKO', 'biurko'),
    (r'[ŁL]AWA', 'lawa'),
    (r'ST[ÓO][ŁL]', 'stol'),
    (r'SZAFA', 'szafa'),
    (r'P[ÓO][ŁL]K[AI]', 'polka'),
    (r'SZAFKA', 'szafka'),
]
ACCESSORY = re.compile(r'UCHWYT|\bLED\b|ZASILAJ|ZASILANIE|PODŚWIETL|OŚWIETL|PROFIL|PANEL ZAGŁOW|PŁOMIEŃ|LISTWA|OCIEKARK|BLENDA|COKÓŁ|BLAT\b|FRONT|PÓŁKI (DO )?SZAF|PÓŁKI SZAFA|KORPUS|NADSTAWK|STELAŻ|MATERAC', re.I)
KITCHEN_MODULE = re.compile(r'SZAFKA (GÓRNA|DOLNA|DONA)|ZMYWARK', re.I)

def classify(d):
    kind = (d.get('kind') or '').upper()
    name = (d.get('tradeName') or '').upper().split(' KPL')[0]
    if KITCHEN_MODULE.search(kind) or (d.get('purpose') or '').upper() == 'KUCHNIA' and 'ZESTAW' not in kind:
        return None, 'moduł kuchenny (kuchnie tylko jako gotowe zestawy)'
    # rodzaj w karcie bywa wpisany błędnie (np. łóżko jako „PODŚWIETLENIE LED”), więc nazwa handlowa ma pierwszeństwo
    if ACCESSORY.search(name) and not re.search(r'ZESTAW MEBLI', kind):
        return None, 'akcesorium / element (nie mebel)'
    for rx, t in TYPE_RULES:
        if re.search(rx, name):
            return t, None
    if ACCESSORY.search(kind):
        return None, 'akcesorium / element (nie mebel)'
    for rx, t in TYPE_RULES:
        if re.search(rx, kind):
            return t, None
    if re.search(r'KONTENER', name): return 'szafka', None
    if re.search(r'BIELIŹNIARK', name): return 'komoda', None
    if re.search(r'WIESZAK', name): return 'wieszak', None
    return None, 'nierozpoznany rodzaj'

ROOM_RULES = [(r'DZIENN|SALON', 'salon'), (r'JADALN', 'jadalnia'), (r'SYPIALN', 'sypialnia'), (r'M[ŁL]ODZIE', 'pokoj-mlodziezowy'),
              (r'PRZEDPOK|GARDEROB', 'przedpokoj'), (r'BIUR|GABINET', 'biuro'), (r'KUCHN', 'kuchnia')]
def rooms_of(purpose):
    p = (purpose or '').upper()
    return [slug for rx, slug in ROOM_RULES if re.search(rx, p)]

def num(v):
    if isinstance(v, (int, float)): return v
    if isinstance(v, str):
        m = re.search(r'\d+(?:[.,]\d+)?', v)
        if m: return float(m.group().replace(',', '.'))
    return None

def http(u):
    # w kartach MATI linki mają spację przed rozszerzeniem („MI01_a%20.jpg”) — na serwerze plik jest bez niej
    return re.sub(r'(?:%20|\s)+(?=\.\w+$)', '', u.strip()) if isinstance(u, str) and u.strip().startswith('http') else None

# ---------------------------------------------------------------- kolor wariantu
TYPE_WORD = re.compile(r'\s(?:RTV|R?KOMODA|R?KREDENS|WITRYNA|REGA[ŁL]|SZAFKA|SZAFA|TOALETKA|LUSTRO|BIURKO|[ŁL]AWA|ST[ÓO][ŁL]|[ŁL]O[ŻZ]E|[ŁL][ÓO][ŻZ]KO|P[ÓO][ŁL]K[AI]|ZESTAW|KONTENEREK|BIELIŹNIARKA|WIESZAK|KPL)(?=[\s./]|$)')
DECOR = [  # skróty i literówki z kart → pełne nazwy dekorów (kolejność ma znaczenie)
    (r'\bCASHMIR\b', 'KASZMIR'), (r'^D\. ?MAGIA$', 'DĄB MAGIA'), (r'\bD\.? CREM\.? TORRO\b', 'DĄB CREMONA TORRO'),
    (r'\bDĄB CR\.? ?T\.?$', 'DĄB CREMONA TORRO'), (r'\bDĄB CREMONA T\.?$', 'DĄB CREMONA TORRO'),
    (r'\bCREM\.?(?=\s)', 'CREMONA'), (r'\bCR\.(?=\s)', 'CREMONA'),
    (r'(?:P[ÓO][ŁL]KA|KOMODA|WITRYNA)?KPL\.?$', ''), (r'\bPOJ\.?$', ''),
    (r'\bBIA[ŁL] PO[ŁL]YSK$', 'BIAŁY POŁYSK'), (r'\bBIA[ŁL] P\.?$', 'BIAŁY POŁYSK'), (r'\bBIA[ŁL]\.?$', 'BIAŁY'), (r'\bPO[ŁL]\.?$', 'POŁYSK'), (r'\bCZARNY P\.?$', 'CZARNY POŁYSK'),
    (r'\bZ[ŁL]\.?$', 'ZŁOTY'), (r'\bBIEL ALP\.?$', 'BIEL ALPEJSKA POŁYSK'), (r'^CZ\.?$', 'CZARNY'),
    (r'^DĄB E\.?$', 'DĄB EVOKE'), (r'^DĄB S\.?$', 'DĄB SONOMA'), (r'^TYP\d\s+', ''),
]
def expand(part):
    part = ' '.join(part.upper().split()).strip(' .,-')
    for rx, full in DECOR:
        part = re.sub(rx, full, part)
    return part.strip()

def parts_of(text):
    return [p for p in (expand(x) for x in re.split(r'\s*[/,]\s*', text or '')) if p]

def name_colors(d, code):
    """Kolory z nazwy handlowej: część między kodem bryły a słowem typu,
    np. „DANTE DT01 DĄB CREMONA T/ANTRACYT MAT RTV KPL” → [DĄB CREMONA TORRO, ANTRACYT MAT]."""
    name = ' '.join((d.get('tradeName') or '').upper().split())
    m = re.match(r'(\D+?)0*(\d+)$', code.upper())
    rx = (re.escape(m.group(1)) + r'-?0*' + m.group(2)) if m else re.escape(code.upper())
    hit = re.search(r'\b' + rx + r'(?:/\d+)?\s+(.*)$', name)
    if not hit:
        return []
    seg = TYPE_WORD.split(' ' + hit.group(1))[0]
    seg = re.sub(r'^\s*\d{2,3}\s+', '', seg)
    return parts_of(seg)

def variant_label(d, collection, code):
    parts = parts_of(d.get('bodyColor')) + parts_of(d.get('frontColor'))
    if not d.get('bodyColor'):
        # brak koloru korpusu (np. DANTE antracyt) — kolory z nazwy handlowej, korpus pierwszy
        named = name_colors(d, code)
        parts = named + [x for x in parts if x not in named]
    elif not d.get('frontColor'):
        # brak koloru frontu (np. łóżka BESTO „DĄB E /BIAŁ P”) — dopisz kolory z nazwy, bez zlepków typu „BIAŁY-CZ”
        parts += [x for x in name_colors(d, code) if x not in parts and '-' not in x and not re.search(r'\d|POJ', x)]
    out = []
    for x in parts:
        if any(x == y or y.startswith(x + ' ') for y in out):
            continue  # „DĄB CRAFT” przy „DĄB CRAFT ZŁOTY”
        out = [y for y in out if not x.startswith(y + ' ')] + [x]
    label = ' / '.join(out)
    return label[:1] + label[1:].lower() if label else None

def slugify(s):
    s = unicodedata.normalize('NFKD', s.replace('ł', 'l').replace('Ł', 'L')).encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')

# ---------------------------------------------------------------- główna pętla
TEMPLATE = None
files = sorted(glob.glob(os.path.join(SRC, '*.xlsx')))
for f in files:
    ws = openpyxl.load_workbook(f, read_only=True).worksheets[0]
    h = [norm(x) for x in next(ws.iter_rows(max_row=1, values_only=True))]
    if h[1] == 'NAZWA HANDLOWA (NAVIREO)' and len(h) >= 134 and 'LINK DO ZDJĘĆ POJEDYNCZYCH BRYŁ 7' not in h:
        TEMPLATE = h
        break

products, excluded, arrangements = {}, [], {}
COLL_FIX = {'UNIWERSALNA': None, 'UNIWERSALNE': None, 'BRYŁA UNIWERSALNA': None, 'LED': None}
for f in files:
    fname = os.path.basename(f)
    file_coll = re.split(r'[ _]', fname)[0].upper()
    for d in read_rows(f):
        coll = (d.get('collection') or file_coll).upper().strip()
        if coll in COLL_FIX or coll.isdigit():
            coll = file_coll  # „UNIWERSALNA” ławy/stoły sprzedawane w ramach kolekcji z pliku
        coll = 'FRESCO' if coll == 'FRESCO' else coll
        code = (d.get('code') or d['symbol'].split()[0]).strip()
        code = re.sub(r'^([A-Z]+)-?(\d+)$', lambda m: f'{m.group(1)}{int(m.group(2)):02d}', code.upper())
        t, why = classify(d)
        if not t:
            excluded.append({'collection': coll, 'symbol': d['symbol'], 'name': d['tradeName'], 'reason': why, 'file': fname})
            continue
        base = code
        extra = {}
        if t == 'kuchnia':
            # zestaw kuchenny: produkt = długość zabudowy (180 … 260 cm); A/B/C to warianty kolorystyczne
            length = int(re.search(r'\b(\d{3})\b', d['symbol']).group(1))
            base = re.sub(r'[A-Z]$', '', code.replace('-', ''))
            code = f'{base} {length}'
            extra = {'kitchenLength': length}
            d['w'] = d.get('w') or length
        elif t == 'lozko':
            # łóżko: produkt = model + szerokość materaca + wersja (z pojemnikiem / bez stelaża); wariant = kolor
            m = re.match(r'([A-Z]+)-?(\d+)(?:/(\d{2,3}))?', d['symbol'].upper())
            base = f'{m.group(1)}{int(m.group(2)):02d}'
            size = int(m.group(3)) if m.group(3) else None
            up = (d['tradeName'] + ' ' + d['symbol']).upper()
            storage = bool(re.search(r'Z POJ|POJEMNIK|\bSIP\b|S\.? ?P\b', up))
            no_slats = 'BEZ STE' in up
            code = base + (f'/{size}' if size else '') + (' P' if storage else '') + (' BS' if no_slats else '')
            extra = {'bed': {'size': size, 'storage': storage, 'slats': not no_slats}}
        key = (coll, code)
        arr = [http(d.get(f'arr{i}')) for i in range(1, 6)]
        arr = [a for a in arr if a]
        arrangements.setdefault(coll, [])
        for a in arr:
            if a not in arrangements[coll]: arrangements[coll].append(a)
        variant = {
            'symbol': d['symbol'], 'ean': str(d['ean']) if d.get('ean') else None,
            'label': variant_label(d, coll, base.split()[0]), 'bodyColor': d.get('bodyColor'), 'frontColor': d.get('frontColor'),
            'photos': [p for p in (http(d.get(f'photo{i}')) for i in range(1, 10)) if p],
            'arrangements': arr, 'description': d.get('description'), 'file': fname,
        }
        p = products.get(key)
        if not p:
            p = products[key] = {
                'collection': coll, 'code': code, 'model': base, **extra, 'type': t, 'kind': d.get('kind'), 'layout': d.get('layout'),
                'rooms': rooms_of(d.get('purpose')), 'style': (d.get('style') or '').lower() or None,
                'dims': {k: num(d.get(k)) for k in ('w', 'h', 'd', 'len') if num(d.get(k))},
                'netKg': num(d.get('netKg')), 'grossKg': num(d.get('grossKg')),
                'delivery72': str(d.get('delivery72') or '').upper() == 'TAK',
                'assembly': d.get('assembly'), 'equipment': d.get('equipment'),
                'drawing': http(d.get('drawing')), 'instruction': http(d.get('instruction')), 'animation': http(d.get('animation')),
                'packages': [], 'materials': {k[2:]: d[k] for k in d if k.startswith('m_') and d[k] is not None},
                'variants': [],
            }
            n = int(num(d.get('packages')) or 0)
            for i in range(1, 9):
                if d.get(f'p{i}_ean') or (i <= n and d.get(f'p{i}_w')):
                    p['packages'].append({k: (str(d.get(f'p{i}_{k}')) if k == 'ean' and d.get(f'p{i}_{k}') else num(d.get(f'p{i}_{k}'))) for k in ('ean', 'w', 'h', 'd', 'kg')})
        if any(v['symbol'] == variant['symbol'] for v in p['variants']):
            continue
        p['variants'].append(variant)
        for r in rooms_of(d.get('purpose')):
            if r not in p['rooms']: p['rooms'].append(r)

# ta sama para kolorów w innej kolejności („Czarny połysk / dąb” i „Dąb / czarny połysk”) → najczęstszy zapis w kolekcji
from collections import Counter
for coll in {p['collection'] for p in products.values()}:
    vs = [v for p in products.values() if p['collection'] == coll for v in p['variants'] if v['label']]
    best = {}
    for lab, n in Counter(v['label'] for v in vs).most_common():
        best.setdefault(frozenset(lab.lower().split(' / ')), lab)
    for v in vs:
        v['label'] = best[frozenset(v['label'].lower().split(' / '))]

out = {
    'source': 'Google Drive: „karty produktu dla białegostoku” (Liveo, udostępnione 28.10.2026), eksport Navireo',
    'imported': date.today().isoformat(),
    'files': [os.path.basename(f) for f in files],
    'arrangements': arrangements,
    'products': sorted(products.values(), key=lambda p: (p['collection'], p['code'])),
    'excluded': excluded,
}
json.dump(out, open(OUT, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
print(f"{len(out['products'])} produktów, {sum(len(p['variants']) for p in out['products'])} wariantów, "
      f"{len({p['collection'] for p in out['products']})} kolekcji, pominięto {len(excluded)} → {OUT}")
